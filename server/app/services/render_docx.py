"""render_docx: OCR result → DOCX via layout_spec (0 LLM calls).

Flow: render_docx_spec() — section/column/indent/tab từ layout spec; không
LLM, không HTML render, không framePr.
"""

from __future__ import annotations

import io
import logging
import re

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.text import (
    WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER,
)
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Inches, Pt, Twips
from docx.table import _Cell

from app.services.layout_spec import BlockSpec, PageSpec, build_page_specs

logger = logging.getLogger("app.services.render_docx")

# ---------------------------------------------------------------- constants
IMAGE_PLACEHOLDER = "[ảnh]"


# ---------------------------------------------------------------- render from IR
def _html_table_rows(text: str) -> list[list[str]]:
    rows = []
    for tr in re.findall(r"<tr[^>]*>(.*?)</tr>", text, re.DOTALL | re.IGNORECASE):
        cells = [
            re.sub(r"<[^>]+>", "", c).strip()
            for c in re.findall(
                r"<t[dh][^>]*>(.*?)</t[dh]>", tr, re.DOTALL | re.IGNORECASE
            )
        ]
        if cells:
            rows.append(cells)
    return rows


def _add_html_table(container, text: str) -> bool:
    rows = _html_table_rows(text)
    if not rows:
        return False
    if isinstance(container, _Cell):
        p0 = container.paragraphs[0]
        if not p0.text and len(container.paragraphs) == 1:
            p0._element.getparent().remove(p0._element)
    cols = max(len(r) for r in rows)
    table = container.add_table(rows=len(rows), cols=cols)
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            table.rows[ri].cells[ci].paragraphs[0].add_run(val)
    return True

def _add_cells_table(container, cells: list[list[dict]]) -> bool:
    """wave-3: cells (kernel parse, có rowspan/colspan/header) → docx table
    python-docx với merge() thật (gridSpan/vMerge) — fidelity 100%, html
    regex path không có span.

    Ceiling: grid ncols = hàng sum colspan lớn nhất (rowspan-tiếp-chéo không
    nới rộng grid) — đủ cho bảng VLM; upgrade path = occupancy scan đầy đủ."""
    rows = len(cells)
    ncols = max(
        (sum(max(1, int(c.get("colspan") or 1)) for c in row) for row in cells),
        default=0,
    )
    if not rows or not ncols:
        return False
    if isinstance(container, _Cell):
        p0 = container.paragraphs[0]
        if not p0.text and len(container.paragraphs) == 1:
            p0._element.getparent().remove(p0._element)
    table = container.add_table(rows=rows, cols=ncols)
    grid: list[list[tuple[int, int] | None]] = [[None] * ncols for _ in range(rows)]
    for ri, row in enumerate(cells):
        ci = 0
        for cell in row:
            while ci < ncols and grid[ri][ci] is not None:
                ci += 1
            if ci >= ncols:
                break
            rs = max(1, int(cell.get("rowspan") or 1))
            cs = max(1, int(cell.get("colspan") or 1))
            for dr in range(rs):
                for dc in range(cs):
                    if ri + dr < rows and ci + dc < ncols:
                        grid[ri + dr][ci + dc] = (ri, ci)
            anchor = table.cell(ri, ci)
            anchor.paragraphs[0].add_run(str(cell.get("text") or ""))
            if rs > 1 or cs > 1:
                end = table.cell(min(ri + rs - 1, rows - 1),
                                 min(ci + cs - 1, ncols - 1))
                anchor.merge(end)
            ci += cs
    return True


_ALIGN_MAP = {
    "left": WD_ALIGN_PARAGRAPH.LEFT,
    "center": WD_ALIGN_PARAGRAPH.CENTER,
    "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
}
_INDENT_INCHES = 0.5


def _all_paragraphs(container):
    """Body + table-cell paragraphs (merged cells repeat in row.cells — idempotent)."""
    for para in container.paragraphs:
        yield para
    for table in container.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from _all_paragraphs(cell)

def _apply_style_overrides(doc, ov: dict) -> None:
    """User style (jobs.layout_style) → every paragraph/run (user beats IR)."""
    font = ov.get("fontFamily")
    size = ov.get("fontSize")
    align = _ALIGN_MAP.get(ov.get("align"))
    indent = ov.get("indent")
    spacing = ov.get("spacing")
    underline = ov.get("underline")
    for para in _all_paragraphs(doc):
        if align is not None:
            para.alignment = align
        if indent is not None:  # 0 → left_indent 0 (explicit clear)
            para.paragraph_format.left_indent = Inches(_INDENT_INCHES * indent)
        if spacing is not None:
            para.paragraph_format.line_spacing = float(spacing)
        for run in para.runs:
            if font is not None:
                run.font.name = font
            if size is not None:
                run.font.size = Pt(size)
            if underline is not None:
                run.underline = underline



# ---------------------------------------------------------------- render from layout spec
# 96 dpi giả định (PaddleOCR-VL trả px theo ảnh gốc) — px → twips/pt.
TWIPS_PER_PX = 1440 / 96
# Bề ngang vùng chữ của trang Letter mặc định (12240 − 1800×2 twips). Toạ độ
# của spec là px trên ảnh scan (640…1500 px) → phải scale về twips theo đúng
# tỉ lệ này, nếu không indent/tab/cột rộng gấp ~1.8 lần và Word bóp méo.
TEXT_WIDTH_TWIPS = 8640.0
PT_PER_PX = 72 / 96
MAX_SPACE_AFTER_PT = 24.0
MAX_SECTIONS = 8            # vượt → linearize 1 cột (file docx vụn, Word khó sửa)
SEC_COL_SPACE_TWIPS = "720"

_SPEC_ALIGN = {
    "left": WD_ALIGN_PARAGRAPH.LEFT,
    "center": WD_ALIGN_PARAGRAPH.CENTER,
    "right": WD_ALIGN_PARAGRAPH.RIGHT,
    "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
}
_SPEC_TAB_ALIGN = {
    "left": WD_TAB_ALIGNMENT.LEFT,
    "center": WD_TAB_ALIGNMENT.CENTER,
    "right": WD_TAB_ALIGNMENT.RIGHT,
}
_SPEC_TAB_LEADER = {
    "none": WD_TAB_LEADER.SPACES,   # val="none" (default OOXML)
    "dots": WD_TAB_LEADER.DOTS,
}


def _scale_for(page_width: float) -> float:
    """px (bề ngang trang scan) → twips (vùng chữ docx)."""
    return TEXT_WIDTH_TWIPS / max(1.0, page_width * TWIPS_PER_PX)


def _set_section_cols(
    section, cols: list, n: int | None = None, scale: float = 1.0
) -> None:
    """w:cols — python-docx không có API; chèn đúng vị trí schema (trước
    w:docGrid) chứ không append thô (Word có thể từ chối file).

    Cột không đều bề ngang (2 cột scan thật: 23-326 | 724-1009) phải khai
    báo từng `w:col w:w` + equalWidth=0, nếu chỉ `w:num="2"` Word chia đều →
    bản docx lệch vị trí so với bản gốc.
    """
    sect_pr = section._sectPr
    for old_el in sect_pr.findall(qn("w:cols")):
        sect_pr.remove(old_el)
    cols_el = OxmlElement("w:cols")
    count = max(1, n if n is not None else len(cols))
    cols_el.set(qn("w:num"), str(count))
    widths = [max(1, round((c.x1 - c.x0) * TWIPS_PER_PX * scale)) for c in cols]
    if count > 1 and len(widths) == count and max(widths) - min(widths) > 40:
        cols_el.set(qn("w:equalWidth"), "0")
        for i, (col, width) in enumerate(zip(cols, widths)):
            col_el = OxmlElement("w:col")
            col_el.set(qn("w:w"), str(width))
            if i < count - 1:      # w:space = khoảng cách SAU cột này
                gap = max(0, round((cols[i + 1].x0 - col.x1) * TWIPS_PER_PX * scale))
                col_el.set(qn("w:space"), str(gap if gap > 0 else SEC_COL_SPACE_TWIPS))
            cols_el.append(col_el)
    else:
        cols_el.set(qn("w:space"), SEC_COL_SPACE_TWIPS)
    sect_pr.insert_element_before(
        cols_el, "w:formProt", "w:vAlign", "w:noEndnote", "w:titlePg",
        "w:textDirection", "w:bidi", "w:rtlGutter", "w:docGrid",
        "w:printerSettings", "w:sectPrChange",
    )

def _add_image_block(
    doc, blk: BlockSpec, crops: dict[str, bytes], scale: float = 1.0
) -> None:
    """Crop bytes cho region_id → ảnh đúng bbox width; thiếu/lỗi → placeholder."""
    data = crops.get(blk.region_id)
    if data:
        try:
            para = doc.add_paragraph()
            para.alignment = _SPEC_ALIGN.get(blk.align, WD_ALIGN_PARAGRAPH.LEFT)
            para.add_run().add_picture(
                io.BytesIO(data), width=Twips(round(blk.w * TWIPS_PER_PX * scale))
            )
            return
        except Exception as exc:  # ảnh hỏng → giữ placeholder, không crash
            logger.info("crop rid=%s lỗi: %r — placeholder", blk.region_id, exc)
    run = doc.add_paragraph().add_run(IMAGE_PLACEHOLDER)
    run.italic = True

def _omml_of(latex: str) -> str | None:
    """LaTeX → OMML (Office Math). None nếu thiếu lib hoặc latex lỗi.

    ponytail: chỉ block công thức ($$...$$) đi đường này; công thức inline
    trong câu vẫn là text. Thêm inline khi thật sự cần.
    """
    try:
        from latex2mathml.converter import convert as _l2m
        from mathml2omml import convert as _m2o
    except ImportError:      # pragma: no cover - dep là main dep
        return None
    try:
        omml = _m2o(_l2m(latex))
    except Exception as exc:  # latex lạ/không parse được → dùng text detex
        logger.info("latex→OMML failed (%r): %r", latex[:40], exc)
        return None
    if not omml or "<m:oMath" not in omml:
        return None
    return omml


def _add_math_block(doc, blk: BlockSpec, omml: str) -> None:
    para = doc.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pf = para.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(min(MAX_SPACE_AFTER_PT, round(blk.space_after_px * PT_PER_PX, 1)))
    xml = omml if "xmlns:m=" in omml else omml.replace(
        "<m:oMath", f"<m:oMath {nsdecls('m', 'w')}", 1
    )
    para._p.append(parse_xml(xml))


def _add_spec_block(
    doc, blk: BlockSpec, crops: dict[str, bytes] | None = None,
    scale: float = 1.0,
) -> None:
    if blk.kind == "image":
        _add_image_block(doc, blk, crops or {}, scale)
        return
    if blk.kind == "table":
        if blk.cells:
            _add_cells_table(doc, blk.cells)
        else:
            _add_html_table(doc, blk.html or blk.text)
        return

    if blk.latex:
        omml = _omml_of(blk.latex)
        if omml:
            _add_math_block(doc, blk, omml)
            return
    para = doc.add_paragraph()
    para.alignment = _SPEC_ALIGN.get(blk.align, WD_ALIGN_PARAGRAPH.LEFT)
    pf = para.paragraph_format
    if blk.align == "left" and blk.indent_px > 0.5:
        pf.left_indent = Twips(round(blk.indent_px * TWIPS_PER_PX * scale))
    if blk.first_line_indent_px > 0.5:
        pf.first_line_indent = Twips(
            round(blk.first_line_indent_px * TWIPS_PER_PX * scale)
        )
    for tab in blk.tabs:
        pf.tab_stops.add_tab_stop(
            Twips(round(tab.pos_px * TWIPS_PER_PX * scale)),
            _SPEC_TAB_ALIGN.get(tab.align, WD_TAB_ALIGNMENT.LEFT),
            _SPEC_TAB_LEADER.get(tab.leader, WD_TAB_LEADER.SPACES),
        )
    pf.space_before = Pt(0)   # Word cộng space_before + space_after
    pf.space_after = Pt(
        min(MAX_SPACE_AFTER_PT, round(blk.space_after_px * PT_PER_PX, 1))
    )
    # "\n" → <w:br/>, "\t" → <w:tab/> (python-docx Run.text xử lý sẵn)
    para.add_run("\n".join(blk.lines or [blk.text]))


def _group_pages(pages: list[PageSpec]) -> list[list[PageSpec]]:
    """Trang liên tiếp cùng số cột → 1 section (linearize nếu quá nhiều)."""
    groups: list[list[PageSpec]] = []
    for pg in pages:
        if groups and len(groups[-1][0].cols) == len(pg.cols):
            groups[-1].append(pg)
        else:
            groups.append([pg])
    if len(groups) > MAX_SECTIONS:
        logger.warning(
            "docx spec: %d section > %d — linearize 1 cột", len(groups), MAX_SECTIONS
        )
        return [pages] if pages else []
    return groups


def render_docx_spec(
    result: dict,
    style_overrides: dict | None = None,
    crops: dict[str, bytes] | None = None,
) -> bytes:
    """OCR result → DOCX qua layout spec (không LLM, không HTML, không framePr).

    crops: {region_id: bytes} — block image có id trong crops → nhúng ảnh
    (width = bbox width px → twips); thiếu key/ảnh lỗi → placeholder [ảnh].
    """
    pages = build_page_specs(result)
    doc = Document()
    for gi, group in enumerate(_group_pages(pages)):
        if gi:
            doc.add_section(WD_SECTION_START.CONTINUOUS)
        _set_section_cols(
            doc.sections[-1], group[0].cols, scale=_scale_for(group[0].width)
        )
        for pg in group:
            page_scale = _scale_for(pg.width)
            for blk in pg.blocks:
                _add_spec_block(doc, blk, crops, page_scale)
    if style_overrides:
        _apply_style_overrides(doc, style_overrides)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()

async def _load_crops(result: dict, crop_loader) -> dict[str, bytes] | None:
    """Gom crop bytes cho mọi block image (gọi loader TUẦN TỰ, lỗi/None → bỏ)."""
    if crop_loader is None:
        return None
    crops: dict[str, bytes] = {}
    for pg in build_page_specs(result):
        for blk in pg.blocks:
            if blk.kind != "image" or blk.region_id in crops:
                continue
            try:
                data = await crop_loader(blk.region_id)
            except Exception as exc:  # partial-fail: crop thiếu → placeholder
                logger.info("crop loader rid=%s failed: %r", blk.region_id, exc)
                continue
            if data:
                crops[blk.region_id] = data
    return crops

async def build_docx(
    result: dict,
    use_llm: bool = True,
    style_overrides: dict | None = None,
    ir_sink: dict | None = None,
    crop_loader=None,
) -> bytes:
    """Build DOCX from OCR result — layout_spec only (0 LLM call).

    crop_loader: async (rid) → bytes | None — wire crop on-demand (batch
    artifact → get_or_make_crop + storage.get); None/lỗi → placeholder.
    ponytail: use_llm/ir_sink giữ cho tương thích test; xoá khi test cập nhật.
    style_overrides (jobs.layout_style) áp sau cùng.
    """
    crops = await _load_crops(result, crop_loader)
    return render_docx_spec(result, style_overrides, crops)
