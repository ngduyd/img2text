"""Plan D2: LLM layout plan → python-docx DOCX export (position-keeping).

Flow: OCR result blocks (id/type/bbox/text≤200) + page dims → POST
{LLM_BASE_URL}/chat/completions (model, temperature 0,
response_format json_object; one retry — a rejected response_format is
dropped on retry) → JSON plan {"columns": "single"|"two",
"blocks": [{"id", "col": "full"|"left"|"right",
"style": "title"|"h1"|"h2"|"body"|"table"|"caption"}]} → python-docx
render (two-column = borderless table, cell margin 0; page break between
pages). LLM unreachable / plan unparsable raises LLMLayoutError — the
export route maps it to 502 and never caches the failure.
"""

import io
import json
import logging
import re
import time

import httpx
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from docx.table import _Cell

from app.core.config import settings
from app.core.db import SessionLocal
from app.services.runtime_settings import get_runtime_setting

logger = logging.getLogger(__name__)

TEXT_TRUNC_CHARS = 200
STYLE_LEVELS = {"title": 0, "h1": 1, "h2": 2}
STYLE_NAMES = {0: "Title", 1: "Heading 1", 2: "Heading 2"}
VALID_COLS = ("full", "left", "right")
VALID_STYLES = ("title", "h1", "h2", "body", "table", "caption")
IMAGE_PLACEHOLDER = "[ảnh]"
PROMPT_INSTRUCTION = (
    'Bạn là chuyên gia bố cục Word. INPUT là JSON {"pages": [{{width, height}}], '
    '"blocks": [{{id, page, type, bbox, text}}]} từ OCR. Trả về JSON với khóa '
    '"columns" ("single" hoặc "two") và "blocks": mảng {{id, col, style}} cho '
    'MỌI id trong input; col ∈ "full"|"left"|"right", style ∈ '
    '"title"|"h1"|"h2"|"body"|"table"|"caption". Chỉ trả về JSON, không kèm '
    "gì khác."
)


class LLMLayoutError(Exception):
    """LLM unreachable / rejected / unparsable plan (route maps to 502)."""


def _extract(result: dict) -> tuple[list[dict], list[tuple[int, int]]]:
    """Flat blocks with unique id + page index, and per-page (width, height)."""
    blocks: list[dict] = []
    dims: list[tuple[int, int]] = []
    seen: set[str] = set()
    for pi, page in enumerate(result.get("pages", [])):
        dims.append((page.get("width") or 0, page.get("height") or 0))
        for region in page.get("regions", []):
            rid = str(region.get("id"))
            if rid in seen:  # region ids restart per page → keep unique
                rid = f"p{pi}-{rid}"
            seen.add(rid)
            image = bool(region.get("placeholder")) or region.get(
                "verdict"
            ) == "image_unextractable"
            blocks.append(
                {
                    "id": rid,
                    "page": pi,
                    "type": "image" if image else "text",
                    "bbox": region.get("bbox"),
                    "text": (region.get("text") or "")[:TEXT_TRUNC_CHARS],
                }
            )
    return blocks, dims


def _prompt(blocks: list[dict], dims: list[tuple[int, int]]) -> str:
    payload = {
        "pages": [{"width": w, "height": h} for w, h in dims],
        "blocks": blocks,
    }
    return PROMPT_INSTRUCTION + "\n" + json.dumps(payload, ensure_ascii=False)


def _parse_plan(content: str) -> dict:
    """Strip optional markdown fence, parse JSON, normalize columns."""
    text = (content or "").strip()
    fence = re.match(r"^```(?:json)?\s*(.*?)\s*```$", text, re.DOTALL)
    if fence:
        text = fence.group(1)
    try:
        plan = json.loads(text)
    except ValueError as exc:
        raise LLMLayoutError(f"plan không phải JSON: {exc}") from exc
    if not isinstance(plan, dict) or not isinstance(plan.get("blocks"), list):
        raise LLMLayoutError("plan thiếu khóa blocks[]")
    plan["columns"] = "two" if plan.get("columns") == "two" else "single"
    plan["blocks"] = [b for b in plan["blocks"] if isinstance(b, dict)]
    return plan


def _log_call(endpoint: str, model: str, n_blocks: int, started: float) -> None:
    """One INFO line per LLM POST — never log prompt/content (security)."""
    logger.info(
        "llm called: fmt=docx-layout model=%s endpoint=%s blocks=%d elapsed=%.1fs",
        model,
        endpoint,
        n_blocks,
        round(time.perf_counter() - started, 1),
    )

async def _plan(blocks: list[dict], dims: list[tuple[int, int]]) -> dict:
    """POST chat/completions; 1 retry (drops response_format when rejected)."""
    model = await get_runtime_setting(SessionLocal, "llm_model", settings.llm_model)
    base_url = await get_runtime_setting(SessionLocal, "llm_base_url", settings.llm_base_url)
    api_key = await get_runtime_setting(SessionLocal, "llm_api_key", settings.llm_api_key)
    payload: dict = {
        "model": model,
        "messages": [{"role": "user", "content": _prompt(blocks, dims)}],
        "temperature": 0,
        "response_format": {"type": "json_object"},
    }
    url = base_url.rstrip("/") + "/chat/completions"
    headers = (
        {"Authorization": f"Bearer {api_key}"}
        if api_key
        else {}  # empty key = no auth header (local compatible)
    )
    async with httpx.AsyncClient(timeout=settings.llm_timeout) as client:
        for attempt in range(2):
            started = time.perf_counter()
            try:
                resp = await client.post(url, json=payload, headers=headers)
            except httpx.HTTPError as exc:
                _log_call(url, model, len(blocks), started)  # failed call: still observable
                logger.warning("llm call failed: fmt=docx-layout error=%s", exc)
                if attempt:
                    raise LLMLayoutError(f"LLM không phản hồi: {exc}") from exc
                continue  # retry 1
            _log_call(url, model, len(blocks), started)
            if resp.status_code < 400:
                try:
                    content = resp.json()["choices"][0]["message"]["content"]
                except (ValueError, KeyError, IndexError, TypeError) as exc:
                    raise LLMLayoutError(f"phản hồi thiếu content: {exc}") from exc
                return _parse_plan(content)
            logger.warning(
                "LLM plan rejected status=%s body=%s",
                resp.status_code,
                resp.text[:200],
            )
            payload.pop("response_format", None)  # server rejected it → retry bare
            if attempt:
                raise LLMLayoutError(f"LLM HTTP {resp.status_code}")
    raise LLMLayoutError("LLM không phản hồi")


# --- Render helpers (python-docx) ---

def _borderless(table, margin_zero: bool = False) -> None:
    """No borders; optionally cell margin 0 (two-column layout table).

    insert_element_before keeps CT_TblPr schema order (Word rejects
    interleaved tblPr children).
    """
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "nil")
        borders.append(el)
    tbl_pr.insert_element_before(
        borders, "w:shd", "w:tblLayout", "w:tblCellMar", "w:tblLook",
        "w:tblCaption", "w:tblDescription",
    )
    if margin_zero:
        mar = OxmlElement("w:tblCellMar")
        for edge in ("top", "left", "bottom", "right"):
            el = OxmlElement(f"w:{edge}")
            el.set(qn("w:w"), "0")
            el.set(qn("w:type"), "dxa")
            mar.append(el)
        tbl_pr.insert_element_before(
            mar, "w:tblLook", "w:tblCaption", "w:tblDescription"
        )

def _cell_paragraph(cell: _Cell):
    """Reuse the cell's default empty paragraph for the first text emit."""
    if len(cell.paragraphs) == 1 and not cell.paragraphs[0].text:
        return cell.paragraphs[0]
    return cell.add_paragraph()

def _html_table_rows(text: str) -> list[list[str]]:
    """Minimal <tr>/<td> regex parse → rows of stripped cell text."""
    rows = []
    for tr in re.findall(r"<tr[^>]*>(.*?)</tr>", text, re.DOTALL | re.IGNORECASE):
        cells = [
            re.sub(r"<[^>]+>", "", c).strip()
            for c in re.findall(
                r"<t[dh][^>]*>(.*?)</t[dh]>", tr, re.DOTALL | re.IGNORECASE
            )
        ]
        if cells:
            rows.append(cells)
    return rows

def _add_html_table(container, text: str) -> bool:
    """Render an HTML table fragment as a simple docx table; False if none."""
    rows = _html_table_rows(text)
    if not rows:
        return False
    if isinstance(container, _Cell):  # nested table needs the empty para gone
        p0 = container.paragraphs[0]
        if not p0.text and len(container.paragraphs) == 1:
            p0._element.getparent().remove(p0._element)
    cols = max(len(r) for r in rows)
    table = container.add_table(rows=len(rows), cols=cols)
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            table.rows[ri].cells[ci].paragraphs[0].add_run(val)
    return True

def _emit(doc, container, block: dict, style: str) -> None:
    """One block → paragraphs / heading / table inside doc or cell."""
    text = block["text"]
    image = block["type"] == "image"
    if image:  # image blocks are never rendered — text placeholder only
        text = IMAGE_PLACEHOLDER
    if not image and style == "table" and _add_html_table(container, text):
        return
    if not image and style in STYLE_LEVELS:
        if hasattr(container, "add_heading"):  # Document body
            container.add_heading(text, level=STYLE_LEVELS[style])
        else:
            para = _cell_paragraph(container)
            para.style = doc.styles[STYLE_NAMES[STYLE_LEVELS[style]]]
            para.add_run(text)
        return
    para = (
        _cell_paragraph(container)
        if isinstance(container, _Cell)
        else container.add_paragraph()
    )
    run = para.add_run(text)
    if style == "caption" and not image:  # italic, small
        run.italic = True
        run.font.size = Pt(9)

def _segments(items: list[tuple[dict, str, str]]) -> list[dict]:
    """Split two-column items into rows-segments split by full-width blocks."""
    segments = [{"left": [], "right": [], "full": None}]
    for block, col, style in items:
        if col == "full":
            segments[-1]["full"] = (block, style)
            segments.append({"left": [], "right": [], "full": None})
        else:
            segments[-1][col].append((block, style))
    return segments

def _render_page(doc, items: list[tuple[dict, str, str]], two_col: bool) -> None:
    if not two_col:
        for block, _col, style in items:
            _emit(doc, doc, block, style)
        return
    table = doc.add_table(rows=0, cols=2)
    _borderless(table, margin_zero=True)
    for seg in _segments(items):
        for i in range(max(len(seg["left"]), len(seg["right"]))):
            row = table.add_row()
            for side, cell in (("left", row.cells[0]), ("right", row.cells[1])):
                if i < len(seg[side]):
                    block, style = seg[side][i]
                    _emit(doc, cell, block, style)
        if seg["full"] is not None:
            block, style = seg["full"]
            row = table.add_row()
            _emit(doc, row.cells[0].merge(row.cells[1]), block, style)
    doc.add_paragraph()  # body must not end inside a table

def _plan_items(plan: dict, blocks: list[dict]) -> list[tuple[dict, str, str]]:
    """Plan order for known ids; missing/unknown ids handled: un-planned
    blocks append at the end (left column / document end — spec fallback)."""
    by_id = {b["id"]: b for b in blocks}
    planned: dict[str, tuple[str, str]] = {}
    order: list[str] = []
    for entry in plan["blocks"]:
        bid = str(entry.get("id", ""))
        if bid not in by_id or bid in planned:
            continue
        col = entry.get("col") if entry.get("col") in VALID_COLS else "full"
        style = entry.get("style") if entry.get("style") in VALID_STYLES else "body"
        planned[bid] = (col, style)
        order.append(bid)
    items: list[tuple[dict, str, str]] = [
        (by_id[bid], *planned[bid]) for bid in order
    ]
    items += [
        (b, "left", "body") for b in blocks if b["id"] not in planned
    ]
    return items

async def build_layout_docx(
    result: dict,
    page_dims: list[tuple[int, int]],
    style_overrides: dict | None = None,
) -> bytes:
    """LLM layout plan → rendered .docx bytes (zip container).

    Plan M: `style_overrides` (jobs.layout_style, typed by the PUT
    route) applied over the finished plan so the user's saved style
    wins the LLM (spec: override trước render = override beats plan).
    """
    blocks, dims = _extract(result)
    if not dims:
        dims = list(page_dims)
    plan = await _plan(blocks, dims) if blocks else {"columns": "single", "blocks": []}
    items = _plan_items(plan, blocks)
    doc = Document()
    page_count = max(len(dims), 1)
    for pi in range(page_count):
        if pi:
            doc.add_page_break()
        page_items = [it for it in items if it[0]["page"] == pi]
        _render_page(doc, page_items, plan["columns"] == "two")
    if style_overrides:
        _apply_style_overrides(doc, style_overrides)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()

_ALIGN_MAP = {
    "left": WD_ALIGN_PARAGRAPH.LEFT,
    "center": WD_ALIGN_PARAGRAPH.CENTER,
    "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
}
# ceiling: 1 indent level = 0.5" left indent (Word's default tab stop);
# upgrade path = mirror web's 1.5rem/level padding when design pins it
_INDENT_INCHES = 0.5

def _all_paragraphs(container):
    """Body + table-cell paragraphs (two-column + nested HTML tables)."""
    for para in container.paragraphs:
        yield para
    for table in container.tables:
        for row in table.rows:
            for cell in row.cells:
                # merged cells repeat in row.cells — pass idempotent, ok
                yield from _all_paragraphs(cell)

def _apply_style_overrides(doc, ov: dict) -> None:
    """Global user style → every paragraph/run (user beats LLM plan)."""
    font = ov.get("fontFamily")
    size = ov.get("fontSize")
    align = _ALIGN_MAP.get(ov.get("align"))
    indent = ov.get("indent")
    spacing = ov.get("spacing")
    underline = ov.get("underline")
    for para in _all_paragraphs(doc):
        if align is not None:
            para.alignment = align
        if indent is not None:  # 0 → left_indent 0 (explicit clear)
            para.paragraph_format.left_indent = Inches(_INDENT_INCHES * indent)
        if spacing is not None:
            para.paragraph_format.line_spacing = float(spacing)
        for run in para.runs:
            if font is not None:
                run.font.name = font
            if size is not None:
                run.font.size = Pt(size)
            if underline is not None:  # underline áp cho MỌI run
                run.underline = underline
