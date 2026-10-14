"""Plan M3-server: docx.html artifact (PUT → cache → export → download).

Covers: PUT /export/docx-layout/html (204 upsert + bytes read-back, oversize
422 exact, queued 409, other-user 403), GET job layout_cache {md, docx,
docx_html} incl. docx = html OR docx-layout, GET export/docx.html/download,
POST export fmt=docx reads docx.html → python-docx assert table/bold/size/
align/underline (no pandoc, no LLM), fallback pandoc path without docx.html.
"""

import io

from docx import Document

from tests.test_jobs_read import other_client  # noqa: F401

HTML = (
    "<h2>Tiêu đề</h2>"
    '<p style="text-align: center">'
    '<span style="font-family: Arial; font-size: 14pt">xin </span>'
    "<b>chào</b><u> bạn</u></p>"
    "<table><tr><td>A</td><td>B</td></tr><tr><td>1</td><td>2</td></tr></table>"
)


async def test_put_docx_html_204_upsert_bytes(
    auth_client, done_job_with_result, storage
):
    key = f"exports/{done_job_with_result}/docx.html"
    await storage.put(key, b"<p>cu</p>")
    r = await auth_client.put(
        f"/api/jobs/{done_job_with_result}/export/docx-layout/html",
        json={"html": HTML},
    )
    assert r.status_code == 204, r.text
    assert (await storage.get(key)).decode("utf-8") == HTML


async def test_put_docx_html_oversize_422(auth_client, done_job_with_result):
    r = await auth_client.put(
        f"/api/jobs/{done_job_with_result}/export/docx-layout/html",
        json={"html": "x" * 1_000_001},
    )
    assert r.status_code == 422
    assert r.json()["detail"] == "Nội dung quá dài."


async def test_put_docx_html_queued_409(auth_client, seed_job):
    r = await auth_client.put(
        f"/api/jobs/{seed_job}/export/docx-layout/html",
        json={"html": "<p>x</p>"},
    )
    assert r.status_code == 409
    assert r.json()["detail"] == "Chưa có kết quả OCR."


async def test_put_docx_html_other_user_403(
    auth_client, other_client, done_job_with_result
):
    r = await other_client.put(
        f"/api/jobs/{done_job_with_result}/export/docx-layout/html",
        json={"html": "<p>x</p>"},
    )
    assert r.status_code == 403


async def test_put_docx_html_invalidates_derived_docx(
    auth_client, done_job_with_result, storage
):
    """Derived exports/{id}/docx + docx-layout must not outlive the html edit."""
    old_docx = f"exports/{done_job_with_result}/docx"
    old_layout = f"exports/{done_job_with_result}/docx-layout"
    await storage.put(old_docx, b"PK-old")
    await storage.put(old_layout, b"PK-layout")
    r = await auth_client.put(
        f"/api/jobs/{done_job_with_result}/export/docx-layout/html",
        json={"html": HTML},
    )
    assert r.status_code == 204
    assert not await storage.exists(old_docx)
    assert not await storage.exists(old_layout)
    assert await storage.exists(f"exports/{done_job_with_result}/docx.html")


async def test_layout_cache_docx_html_flags(auth_client, done_job_with_result, storage):
    r = await auth_client.get(f"/api/jobs/{done_job_with_result}")
    assert r.json()["layout_cache"] == {"md": False, "docx": False, "docx_html": False}
    await storage.put(
        f"exports/{done_job_with_result}/docx.html", HTML.encode("utf-8")
    )
    r = await auth_client.get(f"/api/jobs/{done_job_with_result}")
    # docx = docx.html OR docx-layout
    assert r.json()["layout_cache"] == {"md": False, "docx": True, "docx_html": True}


async def test_download_docx_html(
    auth_client, done_job_with_result, storage
):
    r = await auth_client.get(
        f"/api/jobs/{done_job_with_result}/export/docx.html/download"
    )
    assert r.status_code == 404  # chưa PUT
    await storage.put(
        f"exports/{done_job_with_result}/docx.html", HTML.encode("utf-8")
    )
    r = await auth_client.get(
        f"/api/jobs/{done_job_with_result}/export/docx.html/download"
    )
    assert r.status_code == 200
    assert r.headers["content-type"].startswith("text/html")
    assert r.content.decode("utf-8") == HTML


async def test_export_docx_prefers_docx_html(
    auth_client, done_job_with_result, storage
):
    """fmt=docx with docx.html present → html_to_docx walker (table/font/
    size/align/underline preserved), no pandoc/LLM involved."""
    await storage.put(
        f"exports/{done_job_with_result}/docx.html", HTML.encode("utf-8")
    )
    r = await auth_client.post(
        f"/api/jobs/{done_job_with_result}/export", json={"fmt": "docx"}
    )
    assert r.status_code == 200, r.text
    data = await storage.get(f"exports/{done_job_with_result}/docx")
    doc = Document(io.BytesIO(data))
    texts = [p.text for p in doc.paragraphs]
    assert "Tiêu đề" in texts
    assert "xin chào bạn" in texts
    para = next(p for p in doc.paragraphs if p.text == "xin chào bạn")
    assert para.alignment is not None  # text-align: center → para align
    runs = {r.text: r for r in para.runs}
    assert runs["xin "].font.name == "Arial"
    assert runs["xin "].font.size is not None  # 14pt
    assert runs["chào"].bold is True
    assert runs[" bạn"].underline is True
    table = doc.tables[0]
    assert [c.text for c in table.rows[0].cells] == ["A", "B"]
    assert [c.text for c in table.rows[1].cells] == ["1", "2"]


async def test_export_docx_falls_back_without_docx_html(
    auth_client, done_job_with_result, storage
):
    """No docx.html → pandoc path (existing behavior, docx bytes)."""
    r = await auth_client.post(
        f"/api/jobs/{done_job_with_result}/export", json={"fmt": "docx"}
    )
    assert r.status_code == 200, r.text
    data = await storage.get(f"exports/{done_job_with_result}/docx")
    assert data[:2] == b"PK"


# --- Lane S Task 1: fmt=docx on-demand crop fallback (S1a) ---

async def test_export_docx_on_demand_crop_when_batch_missing(
    auth_client, done_job_with_result, storage, monkeypatch
):
    """Batch crops/{rid}.png missing → get_or_make_crop fallback → img kept
    in docx (spy: bytes reach html_to_docx)."""
    from fastapi import HTTPException
    import app.routers.exports as m
    from tests.conftest import PNG_1PX

    await storage.delete_prefix(f"exports/{done_job_with_result}/docx")
    await storage.put(
        f"exports/{done_job_with_result}/docx.html",
        b'<p>o</p><img src="crops/r0.png">',
    )
    seen: dict = {}
    real = m.html_to_docx
    def spy(html, crops=None):
        seen["crops"] = dict(crops or {})
        return real(html, crops=crops)
    async def fake_crop(job_id, rid, job, result, storage_):
        seen["on_demand"] = rid
        key = f"exports/{job_id}/crops/{rid}.png"
        await storage_.put(key, PNG_1PX)
        return key
    monkeypatch.setattr(m, "html_to_docx", spy)
    monkeypatch.setattr(m, "get_or_make_crop", fake_crop)

    r = await auth_client.post(
        f"/api/jobs/{done_job_with_result}/export", json={"fmt": "docx"}
    )
    assert r.status_code == 200, r.text
    assert seen["on_demand"] == "r0"
    assert seen["crops"]["r0"] == PNG_1PX
    data = await storage.get(f"exports/{done_job_with_result}/docx")
    doc = Document(io.BytesIO(data))
    assert doc.inline_shapes  # ảnh thực sự nhúng (không drop)

async def test_export_docx_crop_http_exception_propagates(
    auth_client, done_job_with_result, storage, monkeypatch
):
    """R27: chỉ swallow KeyError — HTTPException từ get_or_make_crop propagate."""
    from fastapi import HTTPException
    import app.routers.exports as m

    await storage.delete_prefix(f"exports/{done_job_with_result}/docx")
    await storage.put(
        f"exports/{done_job_with_result}/docx.html",
        b'<p>x</p><img src="crops/r0.png">',
    )
    async def boom(job_id, rid, job, result, storage_):
        raise HTTPException(status_code=404, detail="File gốc không tồn tại.")
    monkeypatch.setattr(m, "get_or_make_crop", boom)

    r = await auth_client.post(
        f"/api/jobs/{done_job_with_result}/export", json={"fmt": "docx"}
    )
    assert r.status_code == 404, r.text
