"""render_docx tests: layout_spec renderer (0 LLM) → DOCX."""

import io
import json
from pathlib import Path

from docx import Document

from app.services.render_docx import build_docx

def _ocr_result() -> dict:
    return {
        "pages": [
            {
                "index": 0,
                "width": 800,
                "height": 600,
                "regions": [
                    {"id": "r0", "bbox": [1, 2, 30, 40],
                     "text": "<table><tr><td>A</td><td>B</td></tr>"
                             "<tr><td>1</td><td>2</td></tr></table>",
                     "verdict": "ok", "placeholder": False},
                    {"id": "r1", "bbox": [1, 50, 30, 80],
                     "text": "chú thích ảnh", "verdict": "ok",
                     "placeholder": False},
                ],
            }
        ]
    }


def _zones_result() -> dict:
    """Trang 0 header+content (không footer), trang 1 content-only."""
    return {
        "pages": [
            {
                "index": 0, "width": 800, "height": 600,
                "regions": [
                    {"id": "h0", "bbox": [100, 5, 700, 30],
                     "text": "TRƯỜNG THCS ABC", "verdict": "ok",
                     "placeholder": False},
                    {"id": "c0", "bbox": [50, 60, 750, 100],
                     "text": "Nội dung phần một", "verdict": "ok",
                     "placeholder": False},
                ],
            },
            {
                "index": 1, "width": 800, "height": 600,
                "regions": [
                    {"id": "c1", "bbox": [50, 60, 750, 100],
                     "text": "Nội dung phần hai", "verdict": "ok",
                     "placeholder": False},
                ],
            },
        ]
    }

class TestBuildDocxSpec:
    async def test_build_docx_no_llm(self):
        """use_llm=False → pure heuristic."""
        data = await build_docx(_ocr_result(), use_llm=False)
        assert data[:2] == b"PK"

    async def test_variable_zones_content_only_page_no_invented_header_footer(self):
        """Trang content-only → heuristic không nhân đôi header, không footer thừa."""
        data = await build_docx(_zones_result(), use_llm=False)
        texts = [p.text.strip() for p in Document(io.BytesIO(data)).paragraphs]
        assert texts.count("TRƯỜNG THCS ABC") == 1
        assert "Nội dung phần hai" in texts

    async def test_variable_zones_header_content_no_empty_footer(self):
        """Trang header+content (không footer) → không thêm footer rỗng."""
        result = _zones_result()
        result["pages"].pop(1)
        data = await build_docx(result, use_llm=False)
        texts = [p.text.strip() for p in Document(io.BytesIO(data)).paragraphs]
        assert [t for t in texts if t] == ["TRƯỜNG THCS ABC", "Nội dung phần một"]

    async def test_image_region_without_bbox_keeps_placeholder(self):
        """P2-2: image node thiếu bbox KHÔNG bị drop — marker vẫn render."""
        result = {"pages": [{"index": 0, "width": 800, "height": 600,
            "regions": [
                {"id": "r0", "bbox": [1, 2, 30, 40], "text": "Châu Á",
                 "verdict": "ok", "placeholder": False},
                {"id": "r1", "text": "", "verdict": "image_unextractable",
                 "placeholder": True},  # thiếu bbox
            ]}]}
        data = await build_docx(result, use_llm=False)
        doc = Document(io.BytesIO(data))
        assert any("[ảnh]" in p.text for p in doc.paragraphs)

# --- wave-3 (spec 2026-10-13): region cells-shape → docx table merge span ---

def _cells_result() -> dict:
    cells = [
        [{"text": "A", "rowspan": 2, "colspan": 1, "header": False},
         {"text": "B", "rowspan": 1, "colspan": 2, "header": False}],
        [{"text": "C", "rowspan": 1, "colspan": 1, "header": False}],
    ]
    return {"pages": [{"index": 0, "width": 800, "height": 600, "regions": [
        {"id": "r0", "bbox": [1, 2, 798, 200],
         "text": "A | B | B\nA | C",  # phẳng (kernel)
         "html": "<table><tr><td rowspan=\"2\">A</td><td colspan=\"2\">B</td></tr>"
                 "<tr><td>C</td></tr></table>",
         "cells": cells, "verdict": "ok", "placeholder": False}]}]}

async def test_cells_table_renders_gridspan_vmerge():
    """wave-3: cells-shape → docx table python-docx merge() thật —
    document.xml chứa gridSpan (colspan) + vMerge (rowspan)."""
    import zipfile

    data = await build_docx(_cells_result(), use_llm=False)
    xml = zipfile.ZipFile(io.BytesIO(data)).read("word/document.xml").decode()
    assert "<w:tbl>" in xml
    assert "gridSpan" in xml and "vMerge" in xml  # B colspan=2, A rowspan=2
    doc = Document(io.BytesIO(data))
    assert len(doc.tables) == 1
    t = doc.tables[0]
    assert t.cell(0, 0).text.strip() == "A"
    assert t.cell(0, 1).text.strip() == "B"
    assert t.cell(1, 0).text.strip() == "A"  # vMerge tiếp tục
    assert t.cell(1, 1).text.strip() == "C"

async def test_html_only_region_keeps_old_path():
    """Region chỉ có html (cũ, không cells) → path regex cũ giữ nguyên."""
    result = {"pages": [{"index": 0, "width": 800, "height": 600, "regions": [
        {"id": "r0", "bbox": [1, 2, 798, 40],
         "text": "<table><tr><td>A</td><td>B</td></tr><tr><td>1</td><td>2</td></tr></table>",
         "verdict": "ok", "placeholder": False}]}]}
    data = await build_docx(result, use_llm=False)
    doc = Document(io.BytesIO(data))
    assert len(doc.tables) == 1
    rows = [[c.text for c in row.cells] for row in doc.tables[0].rows]
    assert rows == [["A", "B"], ["1", "2"]]


# ---------------------------------------------------------------- layout spec (phase 1)
SAMPLES = Path(__file__).parent / "fixtures" / "layout_samples"


def _sample(name: str) -> dict:
    return json.loads((SAMPLES / f"{name}.json").read_text(encoding="utf-8"))


async def test_build_docx_spec_2col_sections_and_cols():
    """2 cột → 1 section liên tục + w:cols/@w:num=2 + indent + tab stop."""
    import zipfile

    data = await build_docx(_sample("2col"), use_llm=False)
    doc = Document(io.BytesIO(data))
    assert len(doc.sections) == 1
    xml = zipfile.ZipFile(io.BytesIO(data)).read("word/document.xml").decode()
    assert '<w:cols w:num="2"' in xml
    indented = [p for p in doc.paragraphs if (p.paragraph_format.left_indent or 0) > 0]
    assert indented and indented[0].text.startswith("Đoạn thụt lề")
    tabbed = [p for p in doc.paragraphs if len(p.paragraph_format.tab_stops) == 1]
    assert tabbed and tabbed[0].text.startswith("Ký tên:")
    assert [p.text for p in doc.paragraphs][:3] == [
        "Cột trái đoạn 1 dòng 1.\nCột trái đoạn 1 dòng 2.",
        "Đoạn thụt lề dòng 1.\nĐoạn thụt lề dòng 2.",
        "Ký tên: .....\tTrần B",
    ]


async def test_build_docx_spec_form_tab_leader_dots():
    """Form → tab stop leader DOTS cho từng hàng nhãn/giá trị."""
    from docx.enum.text import WD_TAB_LEADER

    data = await build_docx(_sample("form"), use_llm=False)
    doc = Document(io.BytesIO(data))
    stops = [list(p.paragraph_format.tab_stops) for p in doc.paragraphs]
    leaders = [s[0].leader for s in stops if s]
    assert leaders == [WD_TAB_LEADER.DOTS] * 3
    assert doc.tables and doc.tables[0].cell(0, 0).text == "STT"


async def test_build_docx_spec_1col_single_section_no_cols_override():
    """1 cột → 1 section, w:cols/@w:num=1, text order giữ nguyên."""
    import zipfile

    data = await build_docx(_sample("1col"), use_llm=False)
    doc = Document(io.BytesIO(data))
    assert len(doc.sections) == 1
    xml = zipfile.ZipFile(io.BytesIO(data)).read("word/document.xml").decode()
    assert '<w:cols w:num="1"' in xml
    texts = [p.text for p in doc.paragraphs]
    assert texts[0] == "CHƯƠNG 1: MỞ ĐẦU"
    assert "[ảnh]" in texts

# ---------------------------------------------------------------- crop embed
import base64
import zipfile

from app.services.render_docx import render_docx_spec

PNG_1PX = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mP8z8DwHwAFAAH/"
    "q842iQAAAABJRU5ErkJggg=="
)

def _media_names(data: bytes) -> list[str]:
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        return [n for n in z.namelist() if n.startswith("word/media/")]

def _image_result() -> dict:
    """1 block image, bbox 200x100 (r1)."""
    return {"pages": [{"index": 0, "width": 800, "height": 600, "regions": [
        {"id": "r1", "bbox": [10, 10, 210, 110], "text": "",
         "verdict": "image_unextractable", "placeholder": True}]}]}

async def test_render_docx_spec_embeds_crop_by_bbox_width():
    """crops có rid → 1 inline image, width = bbox w px → twips, không placeholder."""
    from docx.shared import Twips

    data = render_docx_spec(_image_result(), crops={"r1": PNG_1PX})
    doc = Document(io.BytesIO(data))
    assert len(doc.inline_shapes) == 1
    assert "[ảnh]" not in "\n".join(p.text for p in doc.paragraphs)
    expected = Twips(round(200 * 1440 / 96))  # 3000 twips = 1_905_000 EMU
    assert abs(doc.inline_shapes[0].width - expected) <= expected * 0.05

async def test_render_docx_spec_missing_crop_keeps_placeholder():
    """crops thiếu rid → placeholder [ảnh] còn, không crash."""
    data = render_docx_spec(_image_result(), crops={"other": PNG_1PX})
    doc = Document(io.BytesIO(data))
    assert not doc.inline_shapes
    assert "[ảnh]" in "\n".join(p.text for p in doc.paragraphs)

async def test_build_docx_crop_loader_embeds():
    """build_docx crop_loader trả bytes → ảnh nhúng (media part)."""
    async def loader(rid: str) -> bytes:
        return PNG_1PX

    data = await build_docx(_image_result(), crop_loader=loader)
    assert _media_names(data)
    assert len(Document(io.BytesIO(data)).inline_shapes) == 1

async def test_build_docx_crop_loader_error_keeps_placeholder():
    """loader raise → placeholder, không crash."""
    async def loader(rid: str) -> bytes:
        raise RuntimeError("storage down")

    data = await build_docx(_image_result(), crop_loader=loader)
    assert _media_names(data) == []
    assert "[ảnh]" in "\n".join(p.text for p in Document(io.BytesIO(data)).paragraphs)
